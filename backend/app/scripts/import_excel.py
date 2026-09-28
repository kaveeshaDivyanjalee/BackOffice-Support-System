import os
import openpyxl
from sqlalchemy import create_engine, text
from dotenv import load_dotenv

# Database Models auto-create කරගැනීමට backend app එකෙන් Base import කිරීම
try:
    from app.db.database import Base  # හෝ models තියෙන path එක අනුව
except ImportError:
    Base = None

load_dotenv()

DB_USER = os.getenv("DB_USER", "postgres")
DB_PASSWORD = os.getenv("DB_PASSWORD", "YOUR_STRONG_PASSWORD")
DB_HOST = os.getenv("DB_HOST", "db")
DB_PORT = os.getenv("DB_PORT", "5432")
DB_NAME = os.getenv("DB_NAME", "slt_db")

DATABASE_URL = f"postgresql+psycopg2://{DB_USER}:{DB_PASSWORD}@{DB_HOST}:{DB_PORT}/{DB_NAME}"

def import_with_openpyxl(excel_path):
    try:
        print(f"Connecting to PostgreSQL at {DB_HOST}:{DB_PORT} using psycopg2...")
        engine = create_engine(DATABASE_URL)

        # 1. FastAPI models හරහා Tables auto-create කිරීම (Table එක නැත්නම්)
        if Base is not None:
            print("Creating tables if they don't exist...")
            Base.metadata.create_all(bind=engine)

        print(f"Loading workbook: {excel_path}...")
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet = wb.active

        headers = [str(cell.value).strip().lower() for cell in sheet[1] if cell.value is not None]

        rows_to_insert = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if any(row):
                row_dict = {}
                for header, val in zip(headers, row):
                    if header == 'customer_id' and val is not None:
                        val = str(val).strip()
                    row_dict[header] = val
                rows_to_insert.append(row_dict)

        print(f"Found {len(rows_to_insert)} rows to insert...")

        columns_str = ", ".join(headers)
        placeholders_str = ", ".join([f":{col}" for col in headers])
        query = text(f"INSERT INTO customer_status ({columns_str}) VALUES ({placeholders_str})")

        with engine.begin() as conn:
            for row in rows_to_insert:
                conn.execute(query, row)

        print("✅ Data successfully imported into PostgreSQL!")

    except Exception as e:
        print(f"❌ Error during import: {e}")

if __name__ == "__main__":
    excel_file = os.path.join(os.path.dirname(__file__), "customers.xlsx")
    import_with_openpyxl(excel_file)